"""Category-aware Excel I/O for batch price lookups: template, parse, export."""

from __future__ import annotations

import io
import logging
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..categories.base import CategoryPlugin
from ..models import BatchRowResult

log = logging.getLogger(__name__)

# Upload guard-rails: an .xlsx is a zip, so the upload size limit alone doesn't
# stop a small file from expanding to gigabytes in memory (zip bomb).
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DATA_ROWS = 500
MAX_COLUMNS = 50

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# Every batch row needs to say where to look, in addition to the category's
# own fields (product_name/storage, or manufacturer/model).
_SITE_FIELD = ("target_website", "Target Website", ("target website", "site", "url", "website"))
_COUNTRY_FIELD = ("country", "Country", ("country",))


def _category_row_columns(category: CategoryPlugin) -> list[tuple[str, str, tuple[str, ...]]]:
    columns = [(f.name, f.label, (f.name.lower(),) + f.aliases) for f in category.input_fields]
    if not category.is_multi_site:
        columns.append(_SITE_FIELD)
    if category.allowed_countries is None:
        columns.append(_COUNTRY_FIELD)
    return columns


def build_template(category: CategoryPlugin) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    columns = _category_row_columns(category)
    for col_idx, (_, label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = max(18, len(label) + 4)
    # One example row to show the expected shape.
    example = category.known_sites[0] if category.known_sites else "example.com"
    for col_idx, (name, _, _) in enumerate(columns, start=1):
        if name == "target_website":
            ws.cell(row=2, column=col_idx, value=example)
        elif name == "country" and category.allowed_countries:
            ws.cell(row=2, column=col_idx, value=category.allowed_countries[0])
    buf = io.BytesIO()
    wb.save(buf)
    log.info("Built upload template for category=%s (%d column(s))", category.key, len(columns))
    return buf.getvalue()


def _check_archive(data: bytes) -> str | None:
    """Reject non-xlsx files and ones that would expand too much when opened."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        return "Could not read the Excel file. Upload a .xlsx workbook."
    if total > MAX_UNCOMPRESSED_BYTES:
        log.warning("Workbook rejected: expands to %d bytes (limit %d)", total, MAX_UNCOMPRESSED_BYTES)
        return "The Excel file is too large once opened. Split it into smaller files."
    return None


def parse_workbook(data: bytes, category: CategoryPlugin) -> tuple[list[str], list[dict[str, str]], list[str]]:
    """Returns (headers, rows, errors). Rows are dicts keyed by field name."""
    errors: list[str] = []
    log.info("Parsing uploaded workbook for category=%s (%d bytes)", category.key, len(data))
    archive_error = _check_archive(data)
    if archive_error:
        return [], [], [archive_error]
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        log.warning("Workbook parse failed: could not open file - %s", exc)
        return [], [], ["Could not read the Excel file. Upload a .xlsx workbook."]

    ws: Worksheet = wb.active
    rows_iter = ws.iter_rows(values_only=True, max_col=MAX_COLUMNS)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        log.warning("Workbook parse failed: no header row")
        return [], [], ["The file has no header row."]

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    columns = _category_row_columns(category)

    # Map each expected field to the header column index it was found under.
    field_to_col: dict[str, int] = {}
    for name, label, aliases in columns:
        wanted = {label.lower(), name.lower(), *aliases}
        for idx, header in enumerate(headers):
            if header.lower() in wanted:
                field_to_col[name] = idx
                break

    required_names = [f.name for f in category.input_fields if f.required]
    if not category.is_multi_site:
        required_names.append("target_website")
    missing = [
        next(label for n, label, _ in columns if n == name)
        for name in required_names
        if name not in field_to_col
    ]
    if missing:
        errors.append(f"Missing required column(s): {', '.join(missing)}")
        log.warning("Workbook parse failed: missing column(s) %s", missing)
        return headers, [], errors

    rows: list[dict[str, str]] = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None or str(v).strip() == "" for v in raw_row):
            continue
        if len(rows) >= MAX_DATA_ROWS:
            errors.append(f"The file has more than {MAX_DATA_ROWS} rows. Split it into smaller files.")
            log.warning("Workbook parse failed: more than %d data rows", MAX_DATA_ROWS)
            return headers, [], errors
        row: dict[str, str] = {}
        for name, _, _ in columns:
            idx = field_to_col.get(name)
            value = raw_row[idx] if idx is not None and idx < len(raw_row) else None
            if value is not None:
                row[name] = str(value).strip()
        rows.append(row)

    if not rows:
        errors.append("No data rows found below the header.")
        log.warning("Workbook parse failed: no data rows below the header")
    else:
        log.info("Parsed %d data row(s) from workbook (columns: %s)", len(rows), list(field_to_col))
    return headers, rows, errors


_RESULT_COLUMNS = [
    "Subject", "Status", "Least Price Captured", "Max Price Captured",
    "Current Retail Price", "Currency", "Original Price", "Price Type",
    "Availability", "Confidence", "Link Source", "Reasoning Summary",
    "Citations", "Warnings", "Error",
]


def _format_citations(result) -> str:
    if not result or not result.citations:
        return ""
    parts = []
    for c in result.citations:
        price = f"{c.price} {c.currency or ''}".strip() if c.price is not None else "n/a"
        parts.append(f"[{c.domain}] {price} ({c.price_type or 'unknown'}) - {c.url or 'no link'}")
    return " | ".join(parts)


def _write(ws: Worksheet, row: int, column: int, value) -> Cell:
    """Write a cell, never as a formula: values come from uploads, the model and
    web pages, so a string starting with "=" must stay plain text (prevents
    formula injection such as =HYPERLINK(...) / =WEBSERVICE(...))."""
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and cell.data_type == "f":
        cell.data_type = "s"
    return cell


def export_results(category: CategoryPlugin, rows: list[BatchRowResult]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    input_labels = [f.label for f in category.input_fields]
    if not category.is_multi_site:
        input_labels.append("Target Website")
    if category.allowed_countries is None:
        input_labels.append("Country")
    all_headers = input_labels + _RESULT_COLUMNS

    for col_idx, label in enumerate(all_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = max(16, len(label) + 2)

    input_names = [f.name for f in category.input_fields]
    if not category.is_multi_site:
        input_names.append("target_website")
    if category.allowed_countries is None:
        input_names.append("country")

    for r, row in enumerate(rows, start=2):
        col = 1
        for name in input_names:
            _write(ws, r, col, row.input_fields.get(name, ""))
            col += 1

        result = row.result
        finding = result.finding if result else None
        values = [
            result.subject_label if result else "",
            result.status if result else "error",
            result.least_price if result else None,
            result.max_price if result else None,
            finding.price if finding else None,
            finding.currency if finding else None,
            finding.original_price if finding else None,
            finding.price_type if finding else None,
            finding.availability if finding else None,
            finding.confidence if finding else None,
            finding.product_url if finding else None,
            finding.reasoning if finding else None,
            _format_citations(result),
            "; ".join(result.warnings) if result else "",
            (result.error if result else None) or row.error or "",
        ]
        for value in values:
            _write(ws, r, col, value)
            col += 1

    buf = io.BytesIO()
    wb.save(buf)
    log.info("Exported %d result row(s) to workbook for category=%s", len(rows), category.key)
    return buf.getvalue()
